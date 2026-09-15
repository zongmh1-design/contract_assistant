"""生成文档读取测试所需的确定性小文件。"""

from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


FIXTURE_DIR = Path(__file__).parent


def create_pdf(path: Path, pages: list[str]) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4, invariant=True)
    for page_text in pages:
        pdf.setFont("Helvetica", 12)
        pdf.drawString(72, 770, page_text)
        pdf.showPage()
    pdf.save()


def create_docx(path: Path, include_content: bool) -> None:
    document = Document()
    if include_content:
        document.add_heading("Contract Document Reader Test", level=0)
        document.add_paragraph(
            "This paragraph verifies that DOCX body text is extracted in source order."
        )
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Payment term"
        table.cell(1, 1).text = "Thirty days"
        document.add_paragraph("This paragraph appears after the table.")
    document.save(path)


def create_image(path: Path, image_format: str) -> None:
    image = Image.new("RGB", (320, 120), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 45), "Image contract requires OCR", fill="black")
    image.save(path, format=image_format)


def main() -> None:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    create_pdf(
        FIXTURE_DIR / "text_contract.pdf",
        ["Contract PDF fixture with enough readable text for extraction."],
    )
    create_pdf(
        FIXTURE_DIR / "multi_page_contract.pdf",
        [
            "First page contract text appears before the second page.",
            "Second page contract text preserves the expected page order.",
        ],
    )
    create_pdf(FIXTURE_DIR / "empty_text.pdf", [""])
    (FIXTURE_DIR / "corrupted.pdf").write_bytes(b"this is not a valid PDF")

    create_docx(FIXTURE_DIR / "normal_contract.docx", include_content=True)
    create_docx(FIXTURE_DIR / "empty_contract.docx", include_content=False)
    (FIXTURE_DIR / "corrupted.docx").write_bytes(b"this is not a valid DOCX")

    create_image(FIXTURE_DIR / "contract_image.png", "PNG")
    create_image(FIXTURE_DIR / "contract_image.jpg", "JPEG")


if __name__ == "__main__":
    main()
